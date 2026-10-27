import os
import json
import sys
import re
import threading
import tkinter as tk
from datetime import datetime
from tkinter import ttk, filedialog, messagebox

import pandas as pd
from openpyxl import load_workbook, Workbook

try:
    import arabic_reshaper
    from bidi.algorithm import get_display
    HAS_RTL_FIX = True
except ImportError:
    HAS_RTL_FIX = False

def _asset_path(name):
    """Path to a bundled resource. Works frozen (--onefile) and unfrozen."""
    if getattr(sys, "frozen", False):
        return os.path.join(sys._MEIPASS, name)   
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), name)

# ============== ثابت‌های برنامه ==============
if getattr(sys, "frozen", False):
    APP_DIR = os.path.dirname(sys.executable)
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

JSON_DIR = os.path.join(APP_DIR, "json")

REPORT_FILE          = os.path.join(APP_DIR,   "گزارش نهائی.xlsx")
BASE_FILE            = os.path.join(JSON_DIR,  "base.json")
PRICE_FILE           = os.path.join(JSON_DIR,  "prices.json")
EMBEDDED_PRICES_FILE = os.path.join(JSON_DIR,  "prices.json")
EMBEDDED_BASE_FILE   = os.path.join(JSON_DIR,  "base.json")

os.makedirs(JSON_DIR, exist_ok=True)

# ---------- تعریف استان‌ها  ----------
PROVINCE_CONFIGS = {
    "khuzestan": {
        "label": "خوزستان",
        "file_name": "خوزستان.xlsx",
        "output_file": "گزارش خوزستان.xlsx",
        "source_sheet": "گزارش تجمیعی فروش",
        "output_sheet": "فروش خوزستان",
        "default_province": "خوزستان",
        "columns": {
            "factor":      "شماره فاکتور",
            "date":        "تاریخ",
            "seller":      "نام فروشنده",
            "customer":    "نام مشتری",
            "product":     "نام کالا",
            "quantity":    "تعداد کل",
            "varanegar_code": "کد ورانگر",
        },
        "code_from_name": False,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "east_azerbaijan": {
        "label": "آذربایجان شرقی",
        "file_name": "آذربایجان شرقی.xlsx",
        "output_file": "گزارش آذربایجان شرقی.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش آذربایجان شرقی",
        "default_province": "آذربایجان شرقی",
        "columns": {
            "factor":      "شماره فاکتور",
            "date":        "آخرین تاریخ",
            "seller":      "ویزیتور > نام",
            "customer":    "طرف حساب > نام حقوقی",
            "product":     "کالا > نام",
            "quantity":    "فروش خالص > مقدار",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Khorasan": {
        "label": "خراسان",
        "file_name": "خراسان.xlsx",
        "output_file": "گزارش خراسان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش خراسان",
        "default_province": "خراسان",
        "columns": {
            "factor":      "کدفاکتور",
            "date":        "ناريخ",
            "seller":      "نام ويزيتور",
            "customer":    "نام خريدار",
            "product":     "نام کالا",
            "quantity":    ["تعداد", "مقدار"],
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Mazandaran": {
        "label": "مازندران",
        "file_name": "مازندران.xlsx",
        "output_file": "گزارش مازندران.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش مازندران",
        "default_province": "مازندران",
        "columns": {
            "factor":      "شماره فاکتور",
            "date":        "تاریخ فاکتور",
            "seller":      "نام فروشنده",
            "customer":    "نام مشتری",
            "product":     "نام کالا",
            "quantity":    "ت خالص کل",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Arak": {
        "label": "اراک",
        "file_name": "اراک.xlsx",
        "output_file": "گزارش اراک.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش اراک",
        "default_province": "اراک",
        "columns": {
            "factor":      "شماره فاكتور",
            "date":        "تاريخ فاكتور",
            "seller":      "نام واسطه",
            "customer":    "نام طرف حساب",
            "product":     "نام كالا",
            "quantity":    "مقدار",
            "varanegar_code": None,  
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Ardabil": {
        "label": "اردبیل",
        "file_name": "اردبیل.xlsx",
        "output_file": "گزارش اردبیل.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش اردبیل",
        "default_province": "اردبیل",
        "columns": {
            "factor":      "شماره فاكتور",
            "date":        "تاريخ فاكتور",
            "seller":      "نام واسطه",
            "customer":    "نام طرف حساب",
            "product":     "نام كالا",
            "quantity":    "مقدار",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "west_azerbaijan": {
        "label": "آذربایجان غربی",
        "file_name": "آذربایجان غربی.xlsx",
        "output_file": "گزارش آذربایجان غربی.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش آذربایجان غربی",
        "default_province": "آذربایجان غربی",
        "columns": {
            "factor":      "شماره فاکتور",
            "date":        "تاریخ",
            "seller":      "نام ویزیتور",
            "customer":    "نام طرف حساب",
            "product":     "نام کالا",
            "quantity":    "تعداد کل",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Isfahan": {
        "label": "اصفهان",
        "file_name": "اصفهان.xlsx",
        "output_file": "گزارش اصفهان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش اصفهان",
        "default_province": "اصفهان",
        "columns": {
            "factor":      "شماره",
            "date":        "تاریخ",
            "seller":      "نام فروشنده",
            "customer":    "نام مشتری",
            "product":     "نام کالا",
            "quantity":    "مجموع تعداد",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Alborz": {
        "label": "البرز",
        "file_name": "البرز.xlsx",
        "output_file": "گزارش البرز.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش البرز",
        "default_province": "البرز",
        "columns": {
            "factor":      "شماره فاكتور",
            "date":        "تاريخ فاكتور",
            "seller":      "نام واسطه",
            "customer":    "نام طرف حساب",
            "product":     "نام كالا",
            "quantity":    "مقدار",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Elam": {
        "label": "ایلام",
        "file_name": "ایلام.xlsx",
        "output_file": "گزارش ایلام.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش ایلام",
        "default_province": "ایلام",
        "columns": {
            "factor":      "شماره فاكتور",
            "date":        "تاريخ فاكتور",
            "seller":      "ویزیتور",
            "customer":    "نام طرف حساب",
            "product":     "نام كالا",
            "quantity":    "مقدار",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Semnan": {
        "label": "سمنان",
        "file_name": "سمنان.xlsx",
        "output_file": "گزارش سمنان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش سمنان",
        "default_province": "سمنان",
        "columns": {
            "factor":      "شماره فاكتور",
            "date":        "تاريخ فاكتور",
            "seller":      "نام واسطه",
            "customer":    "نام طرف حساب",
            "product":     "نام كالا",
            "quantity":    "مقدار",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Fars": {
        "label": "فارس",
        "file_name": "فارس.xlsx",
        "output_file": "گزارش فارس.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش فارس",
        "default_province": "فارس",
        "columns": {
            "factor":      "شماره حواله",
            "date":        "تاريخ حواله",
            "seller":      "نام فروشنده",
            "customer":    "نام فروشگاه",
            "product":     "نام كالا",
            "quantity":    "تعداد خالص",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Qazvin": {
        "label": "قزوین",
        "file_name": "قزوین.xlsx",
        "output_file": "گزارش قزوین.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش قزوین",
        "default_province": "قزوین",
        "columns": {
            "factor":      "شماره فاكتور",
            "date":        "تاريخ فاكتور",
            "seller":      "نام واسطه",
            "customer":    "نام طرف حساب",
            "product":     "نام كالا",
            "quantity":    "مقدار",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Qom": {
        "label": "قم",
        "file_name": "قم.xlsx",
        "output_file": "گزارش قم.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش قم",
        "default_province": "قم",
        "columns": {
            "factor":      "ش.فاکتور",
            "date":        "تاريخ فاکتور",
            "seller":      "نام فروشنده",
            "customer":    "نام مشتري",
            "product":     "نام کالا",
            "quantity":    "تعداد کل",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Kurdistan": {
        "label": "کردستان",
        "file_name": "کردستان.xlsx",
        "output_file": "گزارش کردستان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش کردستان",
        "default_province": "کردستان",
        "columns": {
            "factor":      "ش.فاکتور",
            "date":        "تاريخ فاکتور",
            "seller":      "نام فروشنده",
            "customer":    "نام مشتري",
            "product":     "نام کالا",
            "quantity":    "تعداد کل",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Kerman": {
        "label": "کرمان ",
        "file_name": "کرمان.xlsx",
        "output_file": "گزارش کرمان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش کرمان",
        "default_province": "کرمان",
        "columns": {
            "factor":      "سريال",
            "date":        "تاريخ",
            "seller":      "بازارياب",
            "customer":    "شخص",
            "product":     "كالا",
            "quantity":    "تعداد",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Kermanshah": {
        "label": "کرمانشاه ",
        "file_name": "کرمانشاه.xlsx",
        "output_file": "گزارش کرمانشاه.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش کرمانشاه",
        "default_province": "کرمانشاه",
        "columns": {
            "factor":      "شماره فاكتور",
            "date":        "تاريخ فاكتور",
            "seller":      "نام واسطه",
            "customer":    "نام طرف حساب",
            "product":     "نام كالا",
            "quantity":    "مقدار",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Golestan": {
        "label": "گلستان ",
        "file_name": "گلستان.xlsx",
        "output_file": "گزارش گلستان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش گلستان",
        "default_province": "گلستان",
        "columns": {
            "factor":      "فاکتور",
            "date":        "تاریخ",
            "seller":      "ویزیتور",
            "customer":    "نام مشتری",
            "product":     "نام کالا",
            "quantity":    "(تعدادي(جزء",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Gilan": {
        "label": "گیلان ",
        "file_name": "گیلان.xlsx",
        "output_file": "گزارش گیلان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش گیلان",
        "default_province": "گیلان",
        "columns": {
            "factor":      "شماره فاکتور *",
            "date":        "تاريخ فاکتور",
            "seller":      "نام فروشنده",
            "customer":    "نام فروشگاه",
            "product":     "نام کالا",
            "quantity":    "تعداد کل",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Lorestan": {
        "label": "لرستان ",
        "file_name": "لرستان.xlsx",
        "output_file": "گزارش لرستان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش لرستان",
        "default_province": "لرستان",
        "columns": {
            "factor":      "فاکتور",
            "date":        "تاریخ",
            "seller":      "شرح روزنامه",
            "customer":    "نام شخص",
            "product":     "کالا",
            "quantity":    "تعداد",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Hamedan": {
        "label": "همدان ",
        "file_name": "همدان.xlsx",
        "output_file": "گزارش همدان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش همدان",
        "default_province": "همدان",
        "columns": {
            "factor":      "شماره فاكتور",
            "date":        "تاريخ فاكتور",
            "seller":      "نام واسطه",
            "customer":    "نام طرف حساب",
            "product":     "نام كالا",
            "quantity":    "مقدار",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Yazd": {
        "label": "یزد ",
        "file_name": "یزد.xlsx",
        "output_file": "گزارش یزد.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش یزد",
        "default_province": "یزد",
        "columns": {
            "factor":      "شماره فاكتور *",
            "date":        "تاريخ فاكتور",
            "seller":      "نام فروشنده",
            "customer":    "نام فروشگاه",
            "product":     "نام كالا",
            "quantity":    "تعداد كل",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Zanjan": {
        "label": "زنجان ",
        "file_name": "زنجان.xlsx",
        "output_file": "گزارش زنجان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش زنجان",
        "default_province": "زنجان",
        "columns": {
            "factor":      "شماره فاکتور",
            "date":        "تاریخ فاکتور",
            "seller":      "نام فروشنده",
            "customer":    "نام مشتری",
            "product":     "نام کالا",
            "quantity":    "تعداد",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
    "Hormozgan": {
        "label": "هرمزگان ",
        "file_name": "هرمزگان.xlsx",
        "output_file": "گزارش هرمزگان.xlsx",
        "source_sheet": None,
        "output_sheet": "فروش هرمزگان",
        "default_province": "هرمزگان",
        "columns": {
            "factor":      "شماره سند",
            "date":        "تاریخ",
            "seller":      "عامل توزیع",
            "customer":    "نام طرف حساب",
            "product":     "نام کالا",
            "quantity":    "تعداد",
            "varanegar_code": None,
        },
        "code_from_name": True,
        "required": ["factor", "date", "seller", "customer", "product"],
    },
}

def get_report_path(province_key):
    """مسیر فایل گزارش مخصوص استان انتخاب‌شده."""
    cfg = PROVINCE_CONFIGS.get(province_key)
    if cfg and cfg.get("output_file"):
        return os.path.join(APP_DIR, cfg["output_file"])
    return REPORT_FILE

# ---------- ستون‌های خروجی ----------
TGT_FACTOR       = "شماره فاکتور"
TGT_DATE         = "تاریخ"
TGT_SELLER       = "نام فروشنده"
TGT_CUSTOMER     = "نام مشتری"
TGT_QUANTITY     = "تعداد"
TGT_UNIT_PRICE   = "نرخ فروش"
TGT_TOTAL_PRICE  = "مبلغ کل"
TGT_VARANEGAR_CODE  = "کد ورانگر"
TGT_VARANEGAR_NAME  = "نام ورانگر"
TGT_GROUP        = "گروه کالا"
TGT_ALIASES      = "نام‌های معادل"
TGT_PROVINCE     = "استان"

OUTPUT_COLUMNS = [
    TGT_FACTOR, TGT_DATE, TGT_CUSTOMER,
    TGT_VARANEGAR_CODE, TGT_VARANEGAR_NAME,
    TGT_QUANTITY, TGT_TOTAL_PRICE,
    TGT_SELLER, TGT_PROVINCE,
]

BASE_COLUMNS = [TGT_VARANEGAR_NAME, TGT_VARANEGAR_CODE, TGT_GROUP, TGT_ALIASES]

LEGACY_NAME_KEY = "نام کالا"

ALIAS_DISPLAY_SEP = " | "
_ALIAS_SPLIT_RE   = re.compile(r'\s*\|\s*')

_SELLER_TITLES_RE = re.compile(
    r'^(?:[آا]قای|[آا]قا|خانم|بانو|جناب|سرکار)\s+',
    re.UNICODE
)

_SELLER_ROLE_RE = re.compile(
    r'\s+(?:ویزیتور|پایون|سرپرست|توزیع\s+کالا)\b',
    re.UNICODE
)

PRICE_CODE   = "کد کالا"
PRICE_NAME   = "نام کالا"
PRICE_VALUE  = "نرخ فروش"
PRICE_COLUMNS = [PRICE_CODE, PRICE_NAME, PRICE_VALUE]


# ============== بارگذاری دیتابیس ==============
def _load_embedded_list(path):
    try:
        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        return data if isinstance(data, list) else []
    except (FileNotFoundError, json.JSONDecodeError):
        return []

EMBEDDED_PRICES = _load_embedded_list(EMBEDDED_PRICES_FILE)

# ============== توابع کمکی ==============
def _is_na(v):
    if v is None:
        return True
    if isinstance(v, (list, dict, tuple, set)):
        return False
    try:
        return bool(pd.isna(v))
    except (TypeError, ValueError):
        return False

def _str_to_aliases(s):
    if s is None:
        return []
    if isinstance(s, (list, tuple)):
        return [str(a).strip() for a in s if str(a).strip()]
    if not str(s).strip():
        return []
    parts = _ALIAS_SPLIT_RE.split(str(s))
    return [p.strip() for p in parts if p.strip()]

def _aliases_to_str(a):
    if not a:
        return ""
    if isinstance(a, str):
        return ALIAS_DISPLAY_SEP.join(_str_to_aliases(a))
    return ALIAS_DISPLAY_SEP.join(str(x).strip() for x in a if str(x).strip())

# ============== تبدیل‌ها ==============
def _to_numeric_code(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    s = str(value).strip()
    if s == "" or s.lower() == "nan":
        return None
    s_clean = s.replace(",", "").replace(" ", "")
    try:
        if re.fullmatch(r"-?\d+", s_clean):
            return int(s_clean)
        if re.fullmatch(r"-?\d+\.\d+", s_clean):
            return float(s_clean)
    except (ValueError, TypeError):
        pass
    return value

def _to_numeric_quantity(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    s = str(value).strip().replace(",", "")
    if s == "" or s.lower() == "nan":
        return None
    try:
        return int(float(s))
    except (ValueError, TypeError):
        return value

def _to_numeric_price(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    s = str(value).strip().replace(",", "").replace(" ", "")
    if s == "" or s.lower() == "nan":
        return None
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except (ValueError, TypeError):
        return None

def _norm_code(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    s = str(value).strip()
    if s == "" or s.lower() == "nan":
        return ""
    s = s.replace(",", "").replace(" ", "")
    try:
        if re.fullmatch(r"-?\d+", s):
            return str(int(s))
        if re.fullmatch(r"-?\d+\.\d+", s):
            f = float(s)
            return str(int(f)) if f.is_integer() else str(f)
    except (ValueError, TypeError):
        pass
    return s

def _clean_seller_name(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    s = str(value)
    s = s.replace('ي', 'ی').replace('ك', 'ک')

    for ar, en in zip('٠١٢٣٤٥٦٧٨٩', '0123456789'):
        s = s.replace(ar, en)
    for fa, en in zip('۰۱۲۳۴۵۶۷۸۹', '0123456789'):
        s = s.replace(fa, en)

    s = re.sub(r'[0-9]+', '', s)
    s = re.sub(r'[!"#$%&\'()*+,\-./:;<=>?@\[\]\\^_`{|}~؟،؛«»…٪]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    s = re.split(r'\s+[ـ—–‒−]\s*', s, maxsplit=1)[0].strip()

    prev = None
    while prev != s:
        prev = s
        s = _SELLER_TITLES_RE.sub('', s, count=1).strip()

    s = re.sub(r'^[A-Za-z]+\s*', '', s).strip()

    before = s
    s = _SELLER_ROLE_RE.split(s, maxsplit=1)[0].strip()
    had_role_marker = (s != before)

    if had_role_marker:
        words = s.split()
        if len(words) > 2:
            s = ' '.join(words[:2])

    return s if s else None

_DIACRITICS_RE = re.compile(
    r'[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED]',
    re.UNICODE
)


def _clean_customer_name(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    s = str(value)
    s = s.replace('ي', 'ی').replace('ك', 'ک')
    s = s.replace('ة', 'ه').replace('ۀ', 'ه')
    s = s.replace('أ', 'ا').replace('إ', 'ا').replace('آ', 'ا')
    s = _DIACRITICS_RE.sub('', s)
    s = s.replace('\u0640', '')
    for ch in ('\u200e', '\u200f', '\u202a', '\u202b',
               '\u202c', '\u202d', '\u202e', '\ufeff', '\u00a0'):
        s = s.replace(ch, ' ')

    for ar, en in zip('٠١٢٣٤٥٦٧٨٩', '0123456789'):
        s = s.replace(ar, en)
    for fa, en in zip('۰۱۲۳۴۵۶۷۸۹', '0123456789'):
        s = s.replace(fa, en)

    s = re.sub(r'[0-9]+', '', s)
    s = re.sub(r'[!"#$%&\'()*+,\-./:;<=>?@\[\]\\^_`{|}~؟،؛«»…٪]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip()
    s = re.sub(r'[A-Za-z]+', ' ', s)

    return s if s else None


def _clean_date(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, datetime):
        return value

    s = str(value).strip()
    if not s:
        return value

    for ar, en in zip('٠١٢٣٤٥٦٧٨٩', '0123456789'):
        s = s.replace(ar, en)
    for fa, en in zip('۰۱۲۳۴۵۶۷۸۹', '0123456789'):
        s = s.replace(fa, en)

    if re.fullmatch(r'\d{8}\.0+', s):
        s = s.split('.', 1)[0]

    m = re.fullmatch(r'(\d{4})\.(\d{1,2})\.(\d{1,2})(.*)', s)
    if m:
        return f"{m.group(1)}/{m.group(2)}/{m.group(3)}{m.group(4)}"

    m = re.fullmatch(r'(\d{4})(\d{2})(\d{2})', s)
    if m:
        y, mo, d = m.group(1), m.group(2), m.group(3)
        if y.startswith(('13', '14')) and 1 <= int(mo) <= 12 and 1 <= int(d)  <= 31:
            return f"{y}/{mo}/{d}"

    m = re.fullmatch(r'(\d{2})[./](\d{1,2})[./](\d{1,2})(.*)', s)
    if m:
        yy, mo, d, rest = m.group(1), m.group(2), m.group(3), m.group(4)
        if 1 <= int(mo) <= 12 and 1 <= int(d) <= 31:
            return f"14{yy}/{int(mo):02d}/{int(d):02d}{rest}"

    m = re.fullmatch(r'(\d{2})(\d{2})(\d{2})', s)
    if m:
        yy, mo, d = m.group(1), m.group(2), m.group(3)
        if 1 <= int(mo) <= 12 and 1 <= int(d) <= 31:
            return f"14{yy}/{mo}/{d}"

    return value


def _extract_trailing_code_from_name(product_name):
    if not product_name:
        return None
    s = str(product_name).strip()
    m = re.search(r'[\s_\-]+(\d{6,12})\s*$', s)
    if m:
        return m.group(1)
    return None


# ============== نرمال‌سازی ==============
_HIDDEN_CHARS = ['\u200c', '\u200d', '\u200e', '\u200f',
                 '\u202a', '\u202b', '\u202c', '\u202d', '\u202e',
                 '\ufeff', '\u00a0']


def normalize_name(name):
    if not name:
        return ""
    name = str(name).strip()
    for ch in _HIDDEN_CHARS:
        name = name.replace(ch, '')
    name = name.replace('ي', 'ی').replace('ك', 'ک')
    name = name.replace('آ', 'ا').replace('أ', 'ا').replace('إ', 'ا')
    name = name.replace('ة', 'ه').replace('ۀ', 'ه')
    for ar, en in zip('٠١٢٣٤٥٦٧٨٩', '0123456789'):
        name = name.replace(ar, en)
    for fa, en in zip('۰۱۲۳۴۵۶۷۸۹', '0123456789'):
        name = name.replace(fa, en)
    name = name.replace('/', '.').replace('،', '.').replace(',', '.')
    name = name.replace('ادریت', 'ادرویت').replace('ادريت', 'ادرویت')
    name = name.replace('آدرويت', 'ادرویت').replace('آدرویت', 'ادرویت')
    name = name.replace('پرفيوم', 'پرفیوم')
    name = name.replace('پرفشنال', 'پروفشنال')
    name = name.replace('کام پروتکت', 'گام پروتکت')
    name = name.replace('دوتایی', 'دو تایی').replace('دو تايی', 'دو تایی')
    name = name.replace('چهارتایی', 'چهار تایی').replace('چهارتایي', 'چهار تایی').replace('چهار تايی', 'چهار تایی')
    name = re.sub(r'\s*(?:کد|كد)?[\s\-_]*\d{3,12}[a-zA-Z]?\s*$', '', name)
    name = re.sub(r'\s+', '', name)
    name = name.replace('_', '').replace('-', '').replace('(', '').replace(')', '')
    name = name.replace('%', '')
    return name.lower()


def normalize_name_spaced(name):
    if not name:
        return ""
    name = str(name).strip()
    for ch in _HIDDEN_CHARS:
        name = name.replace(ch, '')
    name = name.replace('ي', 'ی').replace('ك', 'ک')
    name = name.replace('آ', 'ا').replace('أ', 'ا').replace('إ', 'ا')
    name = name.replace('ة', 'ه').replace('ۀ', 'ه')
    for ar, en in zip('٠١٢٣٤٥٦٧٨٩', '0123456789'):
        name = name.replace(ar, en)
    for fa, en in zip('۰۱۲۳۴۵۶۷۸۹', '0123456789'):
        name = name.replace(fa, en)
    name = name.replace('ادریت', 'ادرویت').replace('ادريت', 'ادرویت')
    name = name.replace('آدرويت', 'ادرویت').replace('آدرویت', 'ادرویت')
    name = name.replace('پرفيوم', 'پرفیوم')
    name = re.sub(r'[^\w\s\u0600-\u06FF]', ' ', name)
    name = re.sub(r'\s+', ' ', name).strip()
    return name.lower()


# ============== تطبیق فازی ==============
_EQUIV_CLASSES = [
    set('تط'), set('سصث'), set('زذضظ'), set('قغ'), set('حه'),
    set('یي'), set('کكگ'), set('اأإآ'), set('وؤ'), set('یئ'),
]


def _chars_equiv(c1, c2):
    if c1 == c2:
        return True
    for cls in _EQUIV_CLASSES:
        if c1 in cls and c2 in cls:
            return True
    return False


def edit_distance_typo(s1, s2):
    if s1 == s2:
        return 0
    if len(s1) < len(s2):
        s1, s2 = s2, s1
    if not s2:
        return len(s1)
    prev = list(range(len(s2) + 1))
    for i, c1 in enumerate(s1, 1):
        curr = [i]
        for j, c2 in enumerate(s2, 1):
            cost = 0 if _chars_equiv(c1, c2) else 1
            curr.append(min(curr[j - 1] + 1, prev[j] + 1, prev[j - 1] + cost))
        prev = curr
    return prev[-1]


def find_fuzzy_match(norm_product, base_keys, max_dist=1):
    n = len(norm_product)
    candidates = []
    for k in base_keys:
        if abs(len(k) - n) > max_dist:
            continue
        d = edit_distance_typo(norm_product, k)
        if d <= max_dist:
            candidates.append((d, k))
    if not candidates:
        return None
    candidates.sort()
    if len(candidates) > 1 and candidates[0][0] == candidates[1][0]:
        return None
    return candidates[0][1]


def find_fuzzy_match_scored(norm_product, base_keys, max_dist=2, min_ratio=0.92):
    n = len(norm_product)
    if n == 0:
        return None

    candidates = []
    for k in base_keys:
        L = len(k)
        if abs(L - n) > max_dist:
            continue
        d = edit_distance_typo(norm_product, k)
        if d > max_dist:
            continue
        max_len = max(n, L)
        ratio = 1.0 - d / max_len
        if ratio < min_ratio:
            continue
        candidates.append((ratio, d, k))

    if not candidates:
        return None

    candidates.sort(key=lambda t: (-t[0], t[1]))
    best = candidates[0]

    if len(candidates) > 1:
        if abs(candidates[1][0] - best[0]) < 0.005:
            return None

    return best[2]


# ============== تطبیق توکنی ==============
_STOP_TOKENS = {
    "پرفیوم", "پرفيوم", "بادی", "بادي", "بدن", "اسپری", "اسپري", "اسپرى",
    "دئودورانت", "دئودرانت", "دئودورنت", "ديودرانت", "دئوورانت",
    "میل", "میلی", "میليمتری", "میلیمتری", "ml",
    "عدد", "عددی", "عددي", "تایی", "تايی", "تايي",
    "با", "بدون", "و",
    "گرم", "گرمي",
}


def _tokenize(spaced_name):
    if not spaced_name:
        return []
    return re.findall(r'[a-zA-Z\u0600-\u06FF]+|\d+', spaced_name)


def find_token_match(spaced_src, base_token_index, threshold=0.7):
    src_tokens = set(_tokenize(spaced_src))
    if not src_tokens:
        return None
    src_meaningful = {t for t in src_tokens if t not in _STOP_TOKENS and not t.isdigit()}
    src_numbers    = {t for t in src_tokens if t.isdigit()}
    if not src_meaningful:
        return None

    candidates = []
    for spaced_base, base_key, _entry in base_token_index:
        base_tokens = set(_tokenize(spaced_base))
        base_numbers = {t for t in base_tokens if t.isdigit()}
        if src_numbers and base_numbers and src_numbers != base_numbers:
            continue
        base_meaningful = {t for t in base_tokens if t not in _STOP_TOKENS and not t.isdigit()}
        if not base_meaningful:
            continue
        inter = len(src_meaningful & base_meaningful)
        union = len(src_meaningful | base_meaningful)
        if union == 0:
            continue
        score = inter / union
        if score >= threshold:
            candidates.append((score, base_key))

    if not candidates:
        return None
    candidates.sort(reverse=True)
    if len(candidates) > 1 and abs(candidates[0][0] - candidates[1][0]) < 1e-9:
        return None
    return candidates[0][1]


# ============== نرمال‌سازی رکوردها ==============
def _normalize_record(rec):
    if not isinstance(rec, dict):
        return None

    drv_name = str(rec.get(TGT_VARANEGAR_NAME, "") or "").strip()
    if not drv_name:
        drv_name = str(rec.get(LEGACY_NAME_KEY, "") or "").strip()

    aliases = rec.get(TGT_ALIASES, [])
    aliases_str = _aliases_to_str(aliases)

    return {
        TGT_VARANEGAR_NAME: drv_name,
        TGT_VARANEGAR_CODE: str(rec.get(TGT_VARANEGAR_CODE, "") or "").strip(),
        TGT_GROUP:       str(rec.get(TGT_GROUP, "") or "").strip(),
        TGT_ALIASES:     aliases_str,
    }


# ============== دیتابیس کالاها ==============
def load_base_data():
    embedded = _load_embedded_list(BASE_FILE)

    records = []
    for rec in embedded:
        n = _normalize_record(rec)
        if n and n[TGT_VARANEGAR_CODE]:
            records.append(n)

    if os.path.exists(BASE_FILE):
        try:
            with open(BASE_FILE, 'r', encoding='utf-8') as f:
                data = json.load(f)
            if isinstance(data, list):
                for rec in data:
                    n = _normalize_record(rec)
                    if n and n[TGT_VARANEGAR_CODE]:
                        records.append(n)
        except Exception:
            pass

    if not records:
        return pd.DataFrame(columns=BASE_COLUMNS)

    df = pd.DataFrame(records)
    for c in BASE_COLUMNS:
        if c not in df.columns:
            df[c] = ""
    df = df[BASE_COLUMNS]
    df = df[df[TGT_VARANEGAR_CODE].astype(str).str.strip() != ""]
    df = df[df[TGT_VARANEGAR_CODE].astype(str).str.lower() != "nan"]

    merged = []
    seen_index = {}
    for _, row in df.iterrows():
        dname = str(row[TGT_VARANEGAR_NAME]).strip()
        if not dname:
            continue

        new_aliases = _str_to_aliases(row[TGT_ALIASES])

        if dname in seen_index:
            idx = seen_index[dname]
            existing_aliases = _str_to_aliases(merged[idx][TGT_ALIASES])
            all_aliases = list(dict.fromkeys(existing_aliases + new_aliases))
            merged[idx][TGT_ALIASES] = _aliases_to_str(all_aliases)
            if not merged[idx][TGT_GROUP]:
                merged[idx][TGT_GROUP] = str(row[TGT_GROUP]).strip()
        else:
            rec = {c: row[c] for c in BASE_COLUMNS}
            rec[TGT_ALIASES] = _aliases_to_str(new_aliases)
            seen_index[dname] = len(merged)
            merged.append(rec)

    return pd.DataFrame(merged, columns=BASE_COLUMNS)


def save_base_data(df):
    os.makedirs(JSON_DIR, exist_ok=True)
    out = []
    for _, row in df.iterrows():
        rec = {
            TGT_VARANEGAR_NAME: ("" if _is_na(row.get(TGT_VARANEGAR_NAME)) else str(row.get(TGT_VARANEGAR_NAME))),
            TGT_VARANEGAR_CODE: ("" if _is_na(row.get(TGT_VARANEGAR_CODE)) else str(row.get(TGT_VARANEGAR_CODE))),
            TGT_GROUP:       ("" if _is_na(row.get(TGT_GROUP)) else str(row.get(TGT_GROUP))),
            TGT_ALIASES:     _str_to_aliases(row.get(TGT_ALIASES, "")),
        }
        out.append(rec)
    with open(BASE_FILE, 'w', encoding='utf-8') as f:
        json.dump(out, f, ensure_ascii=False, indent=4)


def load_price_data():
    df_embedded = pd.DataFrame(EMBEDDED_PRICES, columns=PRICE_COLUMNS)
    if os.path.exists(PRICE_FILE):
        try:
            with open(PRICE_FILE, 'r', encoding='utf-8') as f:
                data = json.load(f)
            df_user = pd.DataFrame(data, columns=PRICE_COLUMNS)
            df_user = df_user[df_user[PRICE_CODE].astype(str).str.strip() != ""]
            df_user = df_user[df_user[PRICE_CODE].astype(str).str.lower() != "nan"]
            df_merged = pd.concat([df_embedded, df_user], ignore_index=True)
            df_merged = df_merged.drop_duplicates(subset=[PRICE_CODE], keep='last').reset_index(drop=True)
            return df_merged
        except Exception:
            pass
    return df_embedded


def save_price_data(df):
    os.makedirs(JSON_DIR, exist_ok=True)
    data = df[PRICE_COLUMNS].fillna("").to_dict(orient="records")
    with open(PRICE_FILE, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=4)


# ============== تطبیق کد قیمت ==============
def _build_price_lookup(price_dict):
    stripped = {}
    for k, v in price_dict.items():
        stripped[k.lstrip('0') or '0'] = v
    return stripped


def _lookup_price(price_dict, price_stripped, code):
    norm = _norm_code(code)
    if not norm:
        return None
    if norm in price_dict:
        return price_dict[norm]
    st = norm.lstrip('0') or '0'
    if st in price_stripped:
        return price_stripped[st]
    padded = norm.zfill(9)
    if padded in price_dict:
        return price_dict[padded]
    return None


# ============== پردازش اصلی ==============
def process_province(province_key, source_path, report_path, log_cb=None, error_cb=None):
    try:
        if province_key not in PROVINCE_CONFIGS:
            if error_cb: error_cb(f"❌ استان نامعتبر: {province_key}")
            return False, []

        cfg = PROVINCE_CONFIGS[province_key]
        col = cfg["columns"]
        output_sheet = cfg["output_sheet"]
        default_province = cfg["default_province"]

        if not report_path:
            report_path = get_report_path(province_key)

        if log_cb: log_cb(f"📂 خواندن فایل {cfg['label']}.xlsx ...")
        try:
            df_src = pd.read_excel(source_path, sheet_name=None)
        except Exception as e:
            if error_cb: error_cb(f"❌ خطا در باز کردن فایل اکسل: {e}")
            return False, []

        sheet_name = None
        if cfg["source_sheet"] and cfg["source_sheet"] in df_src:
            sheet_name = cfg["source_sheet"]
        else:
            for name in df_src.keys():
                if "تجمیعی" in name and "فروش" in name:
                    sheet_name = name
                    break
            if sheet_name is None and len(df_src) > 0:
                sheet_name = list(df_src.keys())[0]

        if sheet_name is None:
            if error_cb: error_cb("❌ شیت مناسب در فایل پیدا نشد.")
            return False, []

        df = df_src[sheet_name].copy()
        df.columns = [str(c).strip() for c in df.columns]

        def _find_col(name):
            if not name:
                return None
            # accept str OR list/tuple of alternatives
            names = name if isinstance(name, (list, tuple)) else [name]
            cleaned = {str(n).strip() for n in names
                    if n is not None and str(n).strip()}
            if not cleaned:
                return None
            for c in df.columns:
                if str(c).strip() in cleaned:
                    return c
            return None

        col_map = {}
        for key in ["factor", "date", "seller", "customer", "product", "quantity", "varanegar_code"]:
            src_name = col.get(key)
            col_map[key] = _find_col(src_name) if src_name else None

        missing = []
        for key in cfg["required"]:
            if col_map.get(key) is None:
                val = col.get(key) or key
                if isinstance(val, (list, tuple)):
                    val = " یا ".join(str(v) for v in val)
                missing.append(str(val))
        if missing:
            if error_cb: error_cb(f"❌ ستون‌های زیر در فایل {cfg['label']} یافت نشدند:\n{', '.join(missing)}")
            return False, []

        has_quantity = col_map["quantity"] is not None
        has_src_varanegar_code = col_map["varanegar_code"] is not None

        if log_cb: log_cb(f"  - 📄 شیت: {sheet_name} - تعداد رکورد: {len(df)}")
        if has_quantity:
            if log_cb: log_cb(f"  - 📊 ستون «{col_map['quantity']}» یافت شد.")
        if has_src_varanegar_code:
            if log_cb: log_cb(f"  - 📊 ستون «{col_map['varanegar_code']}» یافت شد.")
        elif cfg.get("code_from_name"):
            if log_cb: log_cb("  ℹ️ کد ورانگر از انتهای «نام کالا» استخراج می‌شود.")

        if log_cb: log_cb("🔄 بارگذاری دیتابیس کالاها (Embedded Base) ...")
        df_base = load_base_data()

        base_dict = {}
        base_token_index = []
        alias_index_count = 0

        for _, row in df_base.iterrows():
            dname   = str(row.get(TGT_VARANEGAR_NAME, "") or "").strip()
            aliases = _str_to_aliases(row.get(TGT_ALIASES, ""))

            entry = {
                TGT_VARANEGAR_CODE: row.get(TGT_VARANEGAR_CODE),
                TGT_VARANEGAR_NAME: dname,
                TGT_GROUP:       row.get(TGT_GROUP),
            }

            dn = normalize_name(dname)
            if dn:
                base_dict[dn] = entry
                base_token_index.append((normalize_name_spaced(dname), dn, entry))

            for alias in aliases:
                an = normalize_name(alias)
                if an and an not in base_dict:
                    base_dict[an] = entry
                    alias_index_count += 1
                    base_token_index.append((normalize_name_spaced(alias), an, entry))

        base_keys = list(base_dict.keys())
        if log_cb:
            log_cb(f"  - 📦 {len(df_base)} ردیف، {len(base_dict)} کلید یکتا در دیتابیس "
                   f"(شامل {alias_index_count} نام معادل).")

        if log_cb: log_cb("🔄 بارگذاری دیتابیس قیمت‌ها (Embedded Prices) ...")
        df_price = load_price_data()

        price_dict = {}
        for _, row in df_price.iterrows():
            code = _norm_code(row.get(PRICE_CODE))
            price = _to_numeric_price(row.get(PRICE_VALUE))
            if not code:
                continue
            price_dict[code] = price
        price_stripped = _build_price_lookup(price_dict)
        if log_cb: log_cb(f"  - 💰 {len(price_dict)} کد کالای دارای قیمت یافت شد.")

        match_cache = {}
        fuzzy_used = {}
        token_used = {}
        no_price_map = {}

        def get_entry(norm_product, spaced_product=None):
            if not norm_product:
                return None
            if norm_product in match_cache:
                return match_cache[norm_product]

            if norm_product in base_dict:
                match_cache[norm_product] = base_dict[norm_product]
                return base_dict[norm_product]

            matched_key = find_fuzzy_match(norm_product, base_keys, max_dist=1)
            if matched_key is not None:
                entry = base_dict[matched_key]
                match_cache[norm_product] = entry
                fuzzy_used[norm_product] = matched_key
                return entry

            if len(norm_product) >= 15:
                matched_key = find_fuzzy_match_scored(
                    norm_product, base_keys,
                    max_dist=2, min_ratio=0.92
                )
                if matched_key is not None:
                    entry = base_dict[matched_key]
                    match_cache[norm_product] = entry
                    fuzzy_used[norm_product] = matched_key
                    return entry

            if spaced_product:
                token_key = find_token_match(spaced_product, base_token_index, threshold=0.7)
                if token_key is not None:
                    entry = base_dict[token_key]
                    match_cache[norm_product] = entry
                    token_used[norm_product] = token_key
                    return entry

            match_cache[norm_product] = None
            return None

        result_rows = []
        unmatched_set = set()
        matched_count = 0

        for _, row in df.iterrows():
            product_raw = row[col_map["product"]]
            product = str(product_raw).strip() if pd.notna(product_raw) else ""

            varanegar_code = None
            if has_src_varanegar_code:
                src_code_val = row.get(col_map["varanegar_code"])
                if pd.notna(src_code_val):
                    varanegar_code = src_code_val

            if varanegar_code is None and cfg.get("code_from_name"):
                extracted = _extract_trailing_code_from_name(product)
                if extracted:
                    varanegar_code = extracted

            product_for_match = product
            if varanegar_code is not None:
                code_str = str(varanegar_code).strip()
                candidates_to_strip = {code_str, code_str.lstrip('0'),
                                       code_str.zfill(9), code_str.zfill(10)}
                for cand in candidates_to_strip:
                    if cand and cand in product_for_match:
                        product_for_match = re.sub(
                            r'[\s_\-]*' + re.escape(cand) + r'\s*$',
                            '', product_for_match
                        ).strip()
                        break

            norm_product   = normalize_name(product_for_match)
            spaced_product = normalize_name_spaced(product_for_match)
            entry = get_entry(norm_product, spaced_product)

            if entry is None and product:
                unmatched_set.add(product)
            elif entry is not None:
                matched_count += 1

            if varanegar_code is None and entry:
                varanegar_code = entry.get(TGT_VARANEGAR_CODE)

            qty_raw = row[col_map["quantity"]] if has_quantity else None
            qty = _to_numeric_quantity(qty_raw)

            unit_price = _lookup_price(price_dict, price_stripped, varanegar_code) if varanegar_code is not None else None

            if varanegar_code is not None and _norm_code(varanegar_code) and unit_price is None:
                code_key = str(varanegar_code).strip()
                product_label = entry[TGT_VARANEGAR_NAME] if entry else (product or "(نامشخص)")
                if code_key not in no_price_map:
                    no_price_map[code_key] = {"names": set(), "count": 0}
                no_price_map[code_key]["names"].add(product_label)
                no_price_map[code_key]["count"] += 1

            total_price = None
            if qty is not None and unit_price is not None:
                try:
                    total_price = qty * float(unit_price)
                    if float(total_price).is_integer():
                        total_price = int(total_price)
                except (TypeError, ValueError):
                    total_price = None

            if entry:
                varanegar_name_out = entry[TGT_VARANEGAR_NAME]
                group_out = entry[TGT_GROUP]
            else:
                varanegar_name_out = product or None
                group_out = None

            prov_out = default_province

            result_rows.append({
                TGT_FACTOR:       row[col_map["factor"]],
                TGT_DATE:         _clean_date(row[col_map["date"]]),
                TGT_SELLER:       _clean_seller_name(row[col_map["seller"]]),
                TGT_CUSTOMER:     _clean_customer_name(row[col_map["customer"]]),
                TGT_QUANTITY:     qty,
                TGT_TOTAL_PRICE:  total_price,
                TGT_VARANEGAR_CODE:  varanegar_code,
                TGT_VARANEGAR_NAME:  varanegar_name_out,
                TGT_PROVINCE:     prov_out,
            })

        df_result = pd.DataFrame(result_rows, columns=OUTPUT_COLUMNS)
        unmatched = sorted(unmatched_set)

        if fuzzy_used:
            grouped = {}
            for src_norm, base_key in fuzzy_used.items():
                grouped.setdefault(base_key, []).append(src_norm)
            lines = []
            for base_key, src_list in grouped.items():
                base_name = next(
                    (r[TGT_VARANEGAR_NAME] for _, r in df_base.iterrows()
                     if normalize_name(r[TGT_VARANEGAR_NAME]) == base_key), base_key)
                lines.append(f"   • «{base_name}» ← {len(src_list)} نسخه غلط‌املایی")
            if log_cb:
                log_cb(f"✅ {len(fuzzy_used)} غلط املایی به‌صورت خودکار تصحیح شد:")
                for ln in lines[:20]:
                    log_cb(ln)
                if len(lines) > 20:
                    log_cb(f"   ... و {len(lines) - 20} مورد دیگر")

        if token_used:
            if log_cb:
                log_cb(f"✅ {len(token_used)} نام با تطبیق توکنی به دیتابیس متصل شد:")
                for src_norm, base_key in list(token_used.items())[:10]:
                    log_cb(f"   • «{src_norm}» ← «{base_key}»")
                if len(token_used) > 10:
                    log_cb(f"   ... و {len(token_used) - 10} مورد دیگر")

        if no_price_map:
            if log_cb:
                log_cb("")
                log_cb("═" * 78)
                log_cb(f"⚠️ {len(no_price_map)} کد ورانگر بدون قیمت در دیتابیس قیمت‌ها پیدا شد")
                log_cb("═" * 78)
                log_cb("   کد ورانگر   │ نام ورانگر                                     │ رکورد")
                log_cb("   " + "─" * 74)
                for code, info in no_price_map.items():
                    names_str = "، ".join(sorted(info["names"])) or "(نامشخص)"
                    if len(names_str) > 45:
                        names_str = names_str[:42] + "..."
                    log_cb(f"   {code:<11} │ {names_str:<45} │ {info['count']}")
                log_cb("═" * 78)
                log_cb("")
                log_cb("🔧 راه‌حل: در تب «مدیریت قیمت‌ها» این کدها را با نرخ صحیح اضافه کنید.")

        if unmatched:
            preview = "\n".join(f"  - {p}" for p in unmatched[:15])
            tail = f"\n  ... و {len(unmatched) - 15} مورد دیگر" if len(unmatched) > 15 else ""
            if error_cb: error_cb(f"⚠️ {len(unmatched)} کالا در دیتابیس کالا یافت نشد:\n{preview}{tail}")
        else:
            if log_cb:
                log_cb(f"✅ همه کالاها تطبیق داده شدند ({matched_count} رکورد)")
                if not no_price_map:
                    log_cb(f"✅ همه کدهای ورانگر دارای قیمت هستند ({len(price_dict)} کد)")
                else:
                    log_cb(f"ℹ️ {len(price_dict)} کد دارای قیمت، اما {len(no_price_map)} کد بدون قیمت.")

        actual_report_path = report_path
        if os.path.exists(report_path):
            try:
                with open(report_path, "a+b"):
                    pass
            except PermissionError:
                base_path, ext = os.path.splitext(report_path)
                actual_report_path = f"{base_path}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{ext}"
                if error_cb:
                    error_cb(
                        f"⚠️ فایل «{os.path.basename(report_path)}» در حال حاضر باز است.\n"
                        f"   نتیجه در فایل جایگزین ذخیره می‌شود:\n"
                        f"   {actual_report_path}"
                    )

        try:
            if os.path.exists(actual_report_path):
                wb = load_workbook(actual_report_path)
            else:
                wb = Workbook()
                if "Sheet" in wb.sheetnames:
                    del wb["Sheet"]

            if output_sheet in wb.sheetnames:
                del wb[output_sheet]
            ws = wb.create_sheet(output_sheet)

            ws.append(list(OUTPUT_COLUMNS))
            for _, row in df_result.iterrows():
                out = []
                for h in OUTPUT_COLUMNS:
                    v = row[h]
                    if h == TGT_VARANEGAR_CODE:
                        v = _to_numeric_code(v)
                    elif h == TGT_QUANTITY:
                        v = _to_numeric_quantity(v)
                    elif h == TGT_TOTAL_PRICE:
                        v = _to_numeric_price(v)
                    out.append(v)
                ws.append(out)

            wb.save(actual_report_path)

        except PermissionError as pe:
            if error_cb:
                error_cb(
                    f"❌ دسترسی به فایل گزارش امکان‌پذیر نیست:\n{pe}\n\n"
                    f"لطفاً فایل «{os.path.basename(report_path)}» را در Excel ببندید."
                )
            return False, unmatched

        if log_cb:
            log_cb(f"🎉 عملیات موفق. {len(df_result)} رکورد در شیت «{output_sheet}» ذخیره شد.")
            if actual_report_path != report_path:
                log_cb(f"  - 📂 محل ذخیره جایگزین: {os.path.basename(actual_report_path)}")

        return True, unmatched

    except Exception as e:
        import traceback
        if error_cb: error_cb(f"❌ خطای غیرمنتظره:\n{e}\n{traceback.format_exc()}")
        return False, []


# ============== رابط کاربری ==============
class App:
    def __init__(self, root):
        self.root = root
        self.root.title("پایون توزیع کالا - اتوماسیون چند استانی")
        self.root.geometry("1350x850")
        try:
            self.root.iconbitmap(_asset_path("icon.ico"))
        except Exception:
            pass

        self.province_key = tk.StringVar(value="khuzestan")
        self.source_path = tk.StringVar()

        self.base_data = pd.DataFrame(columns=BASE_COLUMNS)
        self.price_data = pd.DataFrame(columns=PRICE_COLUMNS)
        
        # ذخیره لاگ‌های خام برای کپی صحیح در کلیپ‌بورد
        self._raw_log_messages = []

        style = ttk.Style()
        try: 
            style.theme_use("clam")
            # بهبود فونت‌ها و ظاهر کلی
            style.configure('.', font=('Segoe UI', 9))
            style.configure('TButton', font=('Segoe UI', 9, 'bold'), padding=5)
            style.configure('TNotebook.Tab', font=('Segoe UI', 10, 'bold'), padding=[10, 5])
            style.configure('Treeview', font=('Segoe UI', 9), rowheight=25)
            style.configure('Treeview.Heading', font=('Segoe UI', 9, 'bold'))
        except tk.TclError: 
            pass

        self._build_notebook()
        self.reload_base(silent=True)
        self.reload_prices(silent=True)

    def _current_report_path(self):
        return get_report_path(self._current_province_key())

    def _build_notebook(self):
        nb = ttk.Notebook(self.root)
        nb.pack(fill="both", expand=True, padx=10, pady=10)

        self.tab_process = ttk.Frame(nb)
        nb.add(self.tab_process, text=" ⚙️ پردازش استانی ")
        self._build_process_tab()

        self.tab_base = ttk.Frame(nb)
        nb.add(self.tab_base, text=" 📦 مدیریت دیتابیس کالاها ")
        self._build_base_tab()

        self.tab_price = ttk.Frame(nb)
        nb.add(self.tab_price, text=" 💰 مدیریت قیمت‌ها ")
        self._build_price_tab()

    def _setup_paste_support(self, entry, form_vars, single_value=False):
        def paste():
            try:
                text = entry.clipboard_get()
            except tk.TclError:
                return

            if single_value:
                cleaned = re.sub(r'[\r\n\t]+', ALIAS_DISPLAY_SEP, text.strip())
                cleaned = re.sub(r'\s*\|\s*', ALIAS_DISPLAY_SEP, cleaned)
                try:
                    if entry.selection_present():
                        entry.delete(entry.index("sel.first"), entry.index("sel.last"))
                    entry.insert(entry.index("insert"), cleaned)
                except tk.TclError:
                    pass
                return

            if ("\t" in text) or ("\n" in text) or ("\r" in text):
                parts = [p.strip() for p in re.split(r"[\t\n\r]+", text.strip()) if p.strip() != ""]
                if len(parts) > 1:
                    for (_, var), val in zip(form_vars, parts):
                        var.set(val)
                    return
            try:
                if entry.selection_present():
                    entry.delete(entry.index("sel.first"), entry.index("sel.last"))
                entry.insert(entry.index("insert"), text)
            except tk.TclError:
                pass

        def cut():
            try: entry.event_generate("<<Cut>>")
            except tk.TclError: pass

        def copy():
            try: entry.event_generate("<<Copy>>")
            except tk.TclError: pass

        def select_all():
            entry.selection_range(0, tk.END)
            entry.icursor(tk.END)

        menu = tk.Menu(entry, tearoff=0)
        menu.add_command(label="✂️ برش",        command=cut)
        menu.add_command(label="📋 کپی",         command=copy)
        menu.add_command(label="📥 چسباندن",     command=paste)
        menu.add_separator()
        menu.add_command(label="🔤 انتخاب همه",  command=select_all)

        def show_menu(event):
            try:
                entry.focus_set()
                menu.tk_popup(event.x_root, event.y_root)
            finally:
                menu.grab_release()

        entry.bind("<Button-3>",  show_menu)
        entry.bind("<Control-v>", lambda e: (paste(), "break")[1])
        entry.bind("<Control-V>", lambda e: (paste(), "break")[1])
        entry.bind("<Control-a>", lambda e: (select_all(), "break")[1])
        entry.bind("<Control-A>", lambda e: (select_all(), "break")[1])
        entry.bind("<Button-2>",  show_menu)

    # =================================================================
    #                       تب پردازش
    # =================================================================
    def _build_process_tab(self):
        pf = ttk.LabelFrame(self.tab_process, text=" 📍 انتخاب استان ")
        pf.pack(fill="x", padx=10, pady=(10, 5))
        pr = ttk.Frame(pf); pr.pack(fill="x", padx=10, pady=10)

        ttk.Label(pr, text="استان:").pack(side="right")
        self.province_combo = ttk.Combobox(
            pr,
            textvariable=self.province_key,
            state="readonly",
            width=32,
            values=[f"{k} | {v['label']}" for k, v in PROVINCE_CONFIGS.items()],
        )
        self.province_combo.pack(side="right", padx=10)

        cur = self.province_key.get()
        if cur in PROVINCE_CONFIGS:
            self.province_combo.set(f"{cur} | {PROVINCE_CONFIGS[cur]['label']}")
        self.province_combo.bind("<<ComboboxSelected>>", self._on_province_change)

        f = ttk.LabelFrame(self.tab_process, text=" 📁 فایل ورودی ")
        f.pack(fill="x", padx=10, pady=(5, 5))
        r = ttk.Frame(f); r.pack(fill="x", padx=10, pady=10)
        self.src_label = ttk.Label(r, text=self._source_label_text())
        self.src_label.pack(side="right")
        ttk.Entry(r, textvariable=self.source_path, width=70).pack(side="right", padx=10, fill="x", expand=True)
        ttk.Button(r, text="📂 انتخاب فایل…", command=self._select_file).pack(side="right")

        of = ttk.LabelFrame(self.tab_process, text=" 💾 فایل خروجی ")
        of.pack(fill="x", padx=10, pady=(5, 5))
        orow = ttk.Frame(of); orow.pack(fill="x", padx=10, pady=8)
        self.output_label = ttk.Label(orow, text=self._output_label_text(),
                                      foreground="#0a7a3f")
        self.output_label.pack(side="right")

        bf = ttk.Frame(self.tab_process); bf.pack(fill="x", padx=10, pady=10)
        ttk.Button(bf, text="▶️ پردازش و تولید گزارش", command=self._run_process).pack(side="right")
        ttk.Button(bf, text="📂 باز کردن فایل خروجی", command=self._open_output).pack(side="right", padx=8)

        self.status_var = tk.StringVar(value="آماده.")
        ttk.Label(self.tab_process, textvariable=self.status_var, foreground="#0066cc", font=('Segoe UI', 9, 'bold')).pack(anchor="e", padx=10)

        lf = ttk.LabelFrame(self.tab_process, text=" 📊 گزارش عملیات ")
        lf.pack(fill="both", expand=True, padx=10, pady=10)
        
        # استفاده از فونت Tahoma برای نمایش صحیح‌تر فارسی در Tkinter
        self.log_box = tk.Text(lf, wrap="word", font=("Tahoma", 10), height=15, bg="#f9f9f9")
        sb = ttk.Scrollbar(lf, command=self.log_box.yview)
        self.log_box.configure(yscrollcommand=sb.set)
        self.log_box.pack(side="right", fill="both", expand=True)
        sb.pack(side="left", fill="y")

        self.log_box.tag_configure("rtl", justify="right")
        self.log_box.tag_add("rtl", "1.0", "end")
        self.log_box.bind("<<Modified>>", self._reapply_rtl)
        self.log_box.configure(state="normal")
        self.log_box.bind("<Control-c>", self._on_log_ctrl_c)
        self.log_box.bind("<Control-C>", self._on_log_ctrl_c)
        self.log_box.bind("<Control-a>", self._on_log_ctrl_a)
        self.log_box.bind("<Control-A>", self._on_log_ctrl_a)

        self._log_menu = tk.Menu(self.root, tearoff=0)
        self._log_menu.add_command(label="کپی", command=self._on_log_ctrl_c)
        self._log_menu.add_command(label="انتخاب همه", command=self._on_log_ctrl_a)
        self.log_box.bind("<Button-3>", self._show_log_menu)

    def _source_label_text(self):
        key = self._current_province_key()
        cfg = PROVINCE_CONFIGS.get(key)
        return f"{cfg['file_name']}:" if cfg else "فایل:"

    def _output_label_text(self):
        key = self._current_province_key()
        cfg = PROVINCE_CONFIGS.get(key)
        if cfg:
            return f"📄 {cfg.get('output_file', 'گزارش نهائی.xlsx')}"
        return "📄 گزارش نهائی.xlsx"

    def _current_province_key(self):
        raw = self.province_key.get()
        if "|" in raw:
            raw = raw.split("|", 1)[0].strip()
        return raw

    def _on_province_change(self, event=None):
        raw = self.province_combo.get()
        key = raw.split("|", 1)[0].strip() if "|" in raw else raw.strip()
        if key in PROVINCE_CONFIGS:
            self.province_key.set(key)
        self.src_label.configure(text=self._source_label_text())
        if hasattr(self, "output_label"):
            self.output_label.configure(text=self._output_label_text())
        self.source_path.set("")

    def _on_log_ctrl_c(self, event=None):
        try: self.log_box.event_generate("<<Copy>>")
        except tk.TclError: pass
        return "break"

    def _on_log_ctrl_a(self, event=None):
        self.log_box.tag_add("sel", "1.0", "end-1c")
        self.log_box.mark_set("insert", "1.0")
        self.log_box.see("insert")
        return "break"

    def _show_log_menu(self, event):
        try: self._log_menu.tk_popup(event.x_root, event.y_root)
        finally: self._log_menu.grab_release()

    def _copy_all_log(self):
        text = "\n".join(self._raw_log_messages)
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        self.status_var.set("✅ متن اصلی لاگ کپی شد.")

    def _clear_log(self):
        self.log_box.delete(1.0, "end")
        self._raw_log_messages.clear()

    def _reapply_rtl(self, event=None):
        try:
            self.log_box.tag_add("rtl", "1.0", "end")
            self.log_box.edit_modified(False)
        except tk.TclError:
            pass

    # =================================================================
    #              تب مدیریت دیتابیس کالاها
    # =================================================================
    def _build_base_tab(self):
        form = ttk.LabelFrame(self.tab_base, text=" ➕ افزودن / ویرایش کالا ")
        form.pack(fill="x", padx=10, pady=10)

        self.vars = {}
        fields = [
            (TGT_VARANEGAR_NAME, "نام ورانگر *"),
            (TGT_VARANEGAR_CODE, "کد ورانگر *"),
            (TGT_GROUP,       "گروه کالا"),
        ]

        form_vars = []
        entries   = []
        for i, (key, label) in enumerate(fields):
            ttk.Label(form, text=label).grid(row=i, column=1, padx=8, pady=6, sticky="e")
            v = tk.StringVar(value="")
            self.vars[key] = v
            entry = ttk.Entry(form, textvariable=v, width=70)
            entry.grid(row=i, column=0, padx=8, pady=6, sticky="ew")
            form_vars.append((key, v))
            entries.append(entry)
        form.grid_columnconfigure(0, weight=1)

        for e in entries:
            self._setup_paste_support(e, form_vars)

        aliases_row = len(fields)
        ttk.Label(form, text="نام‌های معادل (با | جدا کنید)").grid(
            row=aliases_row, column=1, padx=8, pady=6, sticky="e")

        aliases_var = tk.StringVar(value="")
        self.vars[TGT_ALIASES] = aliases_var
        self.aliases_entry = ttk.Entry(form, textvariable=aliases_var, width=70)
        self.aliases_entry.grid(row=aliases_row, column=0, padx=8, pady=6, sticky="ew")
        self._setup_paste_support(
            self.aliases_entry, [(TGT_ALIASES, aliases_var)], single_value=True
        )
        self.aliases_entry.bind("<FocusIn>",
                                lambda e: self.aliases_entry.xview_moveto(1.0))

        sub_row = aliases_row + 1
        sub = ttk.Frame(form)
        sub.grid(row=sub_row, column=0, columnspan=2,
                 padx=8, pady=(0, 8), sticky="ew")

        ttk.Label(sub, text="✏️ افزودن نام معادل جدید:").pack(side="right", padx=(0, 6))
        ttk.Button(sub, text="➕ افزودن",
                   command=self._quick_add_alias).pack(side="right", padx=4)

        self.new_alias_var = tk.StringVar()
        self.new_alias_entry = ttk.Entry(sub, textvariable=self.new_alias_var)
        self.new_alias_entry.pack(side="right", fill="x", expand=True, padx=4)
        self.new_alias_entry.bind(
            "<Return>", lambda e: (self._quick_add_alias(), "break")[1]
        )
        self.new_alias_entry.bind(
            "<Escape>", lambda e: (self.new_alias_var.set(""), "break")[1]
        )

        ttk.Button(sub, text="🗑 پاک کردن کل لیست",
                   command=lambda: self.vars[TGT_ALIASES].set("")
                   ).pack(side="left", padx=4)

        btn_row = sub_row + 1
        bf = ttk.Frame(form)
        bf.grid(row=btn_row, column=0, columnspan=2, pady=10)
        ttk.Button(bf, text="➕ افزودن / به‌روزرسانی",
                   command=self._add_product).pack(side="right")
        ttk.Button(bf, text="🧹 پاک‌سازی فرم",
                   command=self._clear_form).pack(side="right", padx=6)
        ttk.Button(bf, text="💾 ذخیره دیتابیس",
                   command=self._save_base).pack(side="right", padx=6)
        ttk.Button(bf, text="🔄 بارگذاری مجدد",
                   command=self.reload_base).pack(side="right", padx=6)
        ttk.Button(bf, text="📥 افزودن کالاهای نامنطبق اخیر",
                   command=self._import_unmatched).pack(side="right", padx=6)

        sf = ttk.Frame(self.tab_base); sf.pack(fill="x", padx=10)
        ttk.Label(sf, text="🔍 جستجو:").pack(side="right")
        self.search_var = tk.StringVar()
        self.search_var.trace("w", self._filter_base)
        ttk.Entry(sf, textvariable=self.search_var, width=40).pack(side="right", padx=8)
        ttk.Button(sf, text="🗑 حذف ردیف انتخابی",
                   command=self._delete_selected).pack(side="right")

        tf = ttk.Frame(self.tab_base)
        tf.pack(fill="both", expand=True, padx=10, pady=10)
        cols = ["#"] + BASE_COLUMNS
        self.tree = ttk.Treeview(tf, columns=cols, show="headings")
        for c in cols:
            self.tree.heading(c, text="ردیف" if c == "#" else c)
            if c == "#":
                self.tree.column(c, width=50, anchor="center")
            elif c == TGT_ALIASES:
                self.tree.column(c, width=380, anchor="w")
            else:
                self.tree.column(c, width=220, anchor="center")
        vsb = ttk.Scrollbar(tf, orient="vertical",   command=self.tree.yview)
        hsb = ttk.Scrollbar(tf, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        tf.grid_rowconfigure(0, weight=1)
        tf.grid_columnconfigure(0, weight=1)
        self.tree.bind("<Double-1>", self._on_base_row_select)

        # بهبود ظاهر جدول با رنگ‌های یک در میان
        self.tree.tag_configure('oddrow', background='#f9f9f9')
        self.tree.tag_configure('evenrow', background='#ffffff')

        self.last_unmatched = []
        self._refresh_base_table()

    def _on_base_row_select(self, event=None):
        sel = self.tree.selection()
        if not sel: return
        vals = self.tree.item(sel[0], "values")
        for i, key in enumerate(BASE_COLUMNS):
            if i + 1 < len(vals):
                self.vars[key].set(vals[i + 1])
        try:
            self.aliases_entry.xview_moveto(1.0)
        except tk.TclError:
            pass

    def _quick_add_alias(self):
        new_alias = self.new_alias_var.get().strip()
        if not new_alias:
            return

        current_str = self.vars[TGT_ALIASES].get().strip()
        current_list = _str_to_aliases(current_str)

        if new_alias in current_list:
            messagebox.showinfo("اطلاع", "این نام معادل قبلاً در لیست وجود دارد.")
            self.new_alias_var.set("")
            return

        current_list.append(new_alias)
        self.vars[TGT_ALIASES].set(_aliases_to_str(current_list))
        self.new_alias_var.set("")

        try:
            self.aliases_entry.xview_moveto(1.0)
        except tk.TclError:
            pass
        self.new_alias_entry.focus_set()

    # =================================================================
    #                تب مدیریت قیمت‌ها
    # =================================================================
    def _build_price_tab(self):
        form = ttk.LabelFrame(self.tab_price, text=" ➕ افزودن / ویرایش قیمت کالا ")
        form.pack(fill="x", padx=10, pady=10)

        self.price_vars = {}
        price_fields = [
            (PRICE_CODE,  "کد کالا *"),
            (PRICE_NAME,  "نام کالا"),
            (PRICE_VALUE, "نرخ فروش *"),
        ]

        form_vars = []
        entries   = []
        for i, (key, label) in enumerate(price_fields):
            ttk.Label(form, text=label).grid(row=i, column=1, padx=8, pady=6, sticky="e")
            v = tk.StringVar()
            self.price_vars[key] = v
            entry = ttk.Entry(form, textvariable=v, width=55)
            entry.grid(row=i, column=0, padx=8, pady=6, sticky="ew")
            form_vars.append((key, v))
            entries.append(entry)
        form.grid_columnconfigure(0, weight=1)

        for e in entries:
            self._setup_paste_support(e, form_vars)

        bf = ttk.Frame(form); bf.grid(row=len(price_fields), column=0, columnspan=2, pady=10)
        ttk.Button(bf, text="➕ افزودن / به‌روزرسانی", command=self._add_price).pack(side="right")
        ttk.Button(bf, text="🧹 پاک‌سازی فرم",          command=self._clear_price_form).pack(side="right", padx=6)
        ttk.Button(bf, text="💾 ذخیره دیتابیس قیمت",    command=self._save_prices).pack(side="right", padx=6)
        ttk.Button(bf, text="🔄 بارگذاری مجدد",         command=self.reload_prices).pack(side="right", padx=6)

        sf = ttk.Frame(self.tab_price); sf.pack(fill="x", padx=10)
        ttk.Label(sf, text="🔍 جستجو:").pack(side="right")
        self.price_search_var = tk.StringVar()
        self.price_search_var.trace("w", self._filter_price)
        ttk.Entry(sf, textvariable=self.price_search_var, width=40).pack(side="right", padx=8)
        ttk.Button(sf, text="🗑 حذف ردیف انتخابی", command=self._delete_price_selected).pack(side="right")

        tf = ttk.Frame(self.tab_price); tf.pack(fill="both", expand=True, padx=10, pady=10)
        cols = ["#"] + PRICE_COLUMNS
        self.price_tree = ttk.Treeview(tf, columns=cols, show="headings")
        for c in cols:
            self.price_tree.heading(c, text="ردیف" if c == "#" else c)
            self.price_tree.column(c, width=200 if c != "#" else 50, anchor="center")
        vsb = ttk.Scrollbar(tf, orient="vertical", command=self.price_tree.yview)
        hsb = ttk.Scrollbar(tf, orient="horizontal", command=self.price_tree.xview)
        self.price_tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)
        self.price_tree.grid(row=0, column=0, sticky="nsew")
        vsb.grid(row=0, column=1, sticky="ns")
        hsb.grid(row=1, column=0, sticky="ew")
        tf.grid_rowconfigure(0, weight=1)
        tf.grid_columnconfigure(0, weight=1)
        self.price_tree.bind("<<TreeviewSelect>>", self._on_price_select)

        # بهبود ظاهر جدول با رنگ‌های یک در میان
        self.price_tree.tag_configure('oddrow', background='#f9f9f9')
        self.price_tree.tag_configure('evenrow', background='#ffffff')

        self._refresh_price_table()

    # =================================================================
    #                     توابع عمومی
    # =================================================================
    def _select_file(self):
        key = self._current_province_key()
        cfg = PROVINCE_CONFIGS.get(key, {})
        path = filedialog.askopenfilename(
            title=f"انتخاب فایل {cfg.get('label', '')}",
            filetypes=[("Excel Files", "*.xlsx *.xls")],
        )
        if path:
            self.source_path.set(path)

    def _log(self, msg):
        # ذخیره پیام اصلی (برای کپی در کلیپ‌بورد)
        self._raw_log_messages.append(msg)
        
        # در ویندوز، Tkinter از Bidi و شکل‌دهی حروف به صورت بومی (Native) پشتیبانی می‌کند.
        # استفاده از کتابخانه‌های جانبی در ویندوز باعث دوباره‌کاری و به‌هم‌ریختگی متن می‌شود.
        if sys.platform.startswith("win"):
            display_text = msg
        elif HAS_RTL_FIX:
            try:
                reshaped_text = arabic_reshaper.reshape(msg)
                display_text = get_display(reshaped_text)
            except Exception:
                display_text = msg
        else:
            display_text = msg

        self.log_box.insert("end", display_text + "\n", "rtl")
        self.log_box.see("end")
        self.root.update_idletasks()

    def _run_process(self):
        key = self._current_province_key()
        if key not in PROVINCE_CONFIGS:
            messagebox.showerror("خطا", "استان انتخاب‌شده نامعتبر است.")
            return
        cfg = PROVINCE_CONFIGS[key]

        src = self.source_path.get().strip()
        if not src:
            messagebox.showwarning("هشدار", f"لطفاً ابتدا فایل {cfg['file_name']} را انتخاب کنید.")
            return
        if not os.path.exists(src):
            messagebox.showerror("خطا", "فایل انتخاب‌شده وجود ندارد.")
            return

        report_path = get_report_path(key)

        self._clear_log()
        self.status_var.set(f"در حال پردازش {cfg['label']} …")
        self.root.update()

        def worker():
            ok, unmatched = process_province(
                key, src, report_path,
                log_cb=lambda m: self.root.after(0, lambda: self._log(m)),
                error_cb=lambda m: self.root.after(0, lambda: self._log(m)),
            )
            self.last_unmatched = unmatched or []
            self.root.after(0, lambda: self.status_var.set(
                f"پایان عملیات {cfg['label']}. " + (f"⚠️ {len(unmatched)} کالای نامنطبق." if unmatched else "✅ همه کالاها تطبیق خوردند.")))
            if ok:
                self.root.after(0, lambda: self.reload_base(silent=True))

        threading.Thread(target=worker, daemon=True).start()

    def _open_output(self):
        report_path = self._current_report_path()
        if not os.path.exists(report_path):
            messagebox.showwarning(
                "هشدار",
                f"هنوز فایل «{os.path.basename(report_path)}» ساخته نشده است."
            )
            return
        try:
            os.startfile(report_path)
        except Exception as e:
            messagebox.showerror("خطا", str(e))

    # ---------------- تب کالاها ----------------
    def _refresh_base_table(self, filter_term=""):
        for item in self.tree.get_children():
            self.tree.delete(item)
        term = filter_term.strip().lower()
        for i, (_, row) in enumerate(self.base_data.iterrows(), 1):
            row_vals = {}
            for c in BASE_COLUMNS:
                v = row.get(c)
                if _is_na(v):
                    v = ""
                if c == TGT_ALIASES:
                    v = _aliases_to_str(v) if v else ""
                row_vals[c] = str(v)
            row_str = " ".join(row_vals.values()).lower()
            if term and term not in row_str: continue
            
            tag = 'oddrow' if i % 2 == 0 else 'evenrow'
            self.tree.insert("", "end", values=[i] + [row_vals[c] for c in BASE_COLUMNS], tags=(tag,))

    def _filter_base(self, *args):
        self._refresh_base_table(self.search_var.get())

    def _add_product(self):
        new_row = {}
        for k, v in self.vars.items():
            val = v.get().strip()
            if k == TGT_ALIASES:
                val = _aliases_to_str(_str_to_aliases(val))
            new_row[k] = val

        if not new_row[TGT_VARANEGAR_NAME] or not new_row[TGT_VARANEGAR_CODE]:
            messagebox.showwarning("هشدار", "وارد کردن «نام ورانگر» و «کد ورانگر» الزامی است.")
            return

        existing = self.base_data[TGT_VARANEGAR_NAME].astype(str).str.strip().tolist()
        if new_row[TGT_VARANEGAR_NAME] in existing:
            if not messagebox.askyesno("تکراری", "این کالا قبلاً ثبت شده. ویرایش کنیم؟"):
                return
            idx = existing.index(new_row[TGT_VARANEGAR_NAME])
            for k, v in new_row.items():
                self.base_data.at[idx, k] = v
            self._log(f"✏️ کالا ویرایش شد: {new_row[TGT_VARANEGAR_NAME]}")
        else:
            self.base_data = pd.concat([self.base_data, pd.DataFrame([new_row])], ignore_index=True)
            self._log(f"➕ کالا اضافه شد: {new_row[TGT_VARANEGAR_NAME]}")
        self._refresh_base_table(self.search_var.get())
        self._clear_form()

    def _clear_form(self):
        for v in self.vars.values():
            v.set("")
        if hasattr(self, "new_alias_var"):
            self.new_alias_var.set("")

    def _save_base(self):
        try:
            save_base_data(self.base_data)
            messagebox.showinfo("موفق", f"دیتابیس کالاها در فایل base.json ذخیره شد:\n{BASE_FILE}")
            self._log("💾 دیتابیس کالاها ذخیره شد.")
        except Exception as e:
            messagebox.showerror("خطا", f"ذخیره ناموفق:\n{e}")

    def reload_base(self, silent=False):
        self.base_data = load_base_data()
        self._refresh_base_table(self.search_var.get() if hasattr(self, "search_var") else "")
        if not silent:
            messagebox.showinfo("موفق", "دیتابیس کالاها بارگذاری مجدد شد.")
        elif hasattr(self, "log_box"):
            self._log(f"📦 دیتابیس کالاها بارگذاری شد ({len(self.base_data)} رکورد).")

    def _delete_selected(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("هشدار", "یک ردیف را انتخاب کنید.")
            return
        if not messagebox.askyesno("تأیید حذف", "حذف شود؟"): return
        idxs = sorted([int(self.tree.item(i, "values")[0]) - 1 for i in sel], reverse=True)
        for idx in idxs:
            if 0 <= idx < len(self.base_data):
                self.base_data = self.base_data.drop(self.base_data.index[idx]).reset_index(drop=True)
        self._refresh_base_table(self.search_var.get())

    def _import_unmatched(self):
        if not getattr(self, "last_unmatched", None):
            messagebox.showinfo("اطلاع", "کالای نامنطبقی موجود نیست.")
            return
        added = 0
        existing = set(self.base_data[TGT_VARANEGAR_NAME].astype(str).str.strip().tolist())
        for p in self.last_unmatched:
            if p in existing: continue
            new_row = {c: "" for c in BASE_COLUMNS}
            new_row[TGT_VARANEGAR_NAME] = p
            new_row[TGT_ALIASES]     = p
            self.base_data = pd.concat([self.base_data, pd.DataFrame([new_row])], ignore_index=True)
            existing.add(p)
            added += 1
        self._refresh_base_table(self.search_var.get())
        messagebox.showinfo("موفق", f"{added} کالا اضافه شد. لطفاً کد ورانگر را پر کرده و ذخیره کنید.")

    # ---------------- تب قیمت‌ها ----------------
    def _refresh_price_table(self, filter_term=""):
        for item in self.price_tree.get_children():
            self.price_tree.delete(item)
        term = filter_term.strip().lower()
        for i, (_, row) in enumerate(self.price_data.iterrows(), 1):
            row_vals = {c: ("" if _is_na(row.get(c)) else str(row.get(c))) for c in PRICE_COLUMNS}
            row_str = " ".join(row_vals.values()).lower()
            if term and term not in row_str: continue
            
            tag = 'oddrow' if i % 2 == 0 else 'evenrow'
            self.price_tree.insert("", "end", values=[i] + [row_vals[c] for c in PRICE_COLUMNS], tags=(tag,))

    def _filter_price(self, *args):
        self._refresh_price_table(self.price_search_var.get())

    def _on_price_select(self, event=None):
        sel = self.price_tree.selection()
        if not sel: return
        vals = self.price_tree.item(sel[0], "values")
        if len(vals) >= 4:
            self.price_vars[PRICE_CODE].set(vals[1])
            self.price_vars[PRICE_NAME].set(vals[2])
            self.price_vars[PRICE_VALUE].set(vals[3])

    def _add_price(self):
        code = self.price_vars[PRICE_CODE].get().strip()
        name = self.price_vars[PRICE_NAME].get().strip()
        val_raw = self.price_vars[PRICE_VALUE].get().strip()
        if not code or not val_raw:
            messagebox.showwarning("هشدار", "وارد کردن «کد کالا» و «نرخ فروش» الزامی است.")
            return
        try:
            val = float(val_raw.replace(",", ""))
            if val.is_integer():
                val = int(val)
        except ValueError:
            messagebox.showerror("خطا", "نرخ فروش باید عدد باشد.")
            return

        new_row = {PRICE_CODE: code, PRICE_NAME: name, PRICE_VALUE: val}
        codes = self.price_data[PRICE_CODE].astype(str).str.strip().tolist()
        if code in codes:
            if not messagebox.askyesno("تکراری", "این کد کالا قبلاً ثبت شده. ویرایش شود؟"):
                return
            idx = codes.index(code)
            for k, v in new_row.items():
                self.price_data.at[idx, k] = v
            self._log(f"✏️ قیمت ویرایش شد: {code} → {val:,}")
        else:
            self.price_data = pd.concat([self.price_data, pd.DataFrame([new_row])], ignore_index=True)
            self._log(f"➕ قیمت اضافه شد: {code} → {val:,}")
        self._refresh_price_table(self.price_search_var.get())
        self._clear_price_form()

    def _clear_price_form(self):
        for v in self.price_vars.values():
            v.set("")

    def _save_prices(self):
        try:
            save_price_data(self.price_data)
            messagebox.showinfo("موفق", f"دیتابیس قیمت‌ها در فایل prices.json ذخیره شد:\n{PRICE_FILE}")
            self._log("💾 دیتابیس قیمت‌ها ذخیره شد.")
        except Exception as e:
            messagebox.showerror("خطا", f"ذخیره ناموفق:\n{e}")

    def reload_prices(self, silent=False):
        self.price_data = load_price_data()
        self._refresh_price_table(self.price_search_var.get() if hasattr(self, "price_search_var") else "")
        if not silent:
            messagebox.showinfo("موفق", "دیتابیس قیمت‌ها بارگذاری مجدد شد.")
        elif hasattr(self, "log_box"):
            self._log(f"💰 دیتابیس قیمت‌ها بارگذاری شد ({len(self.price_data)} رکورد).")

    def _delete_price_selected(self):
        sel = self.price_tree.selection()
        if not sel:
            messagebox.showwarning("هشدار", "یک ردیف را انتخاب کنید.")
            return
        if not messagebox.askyesno("تأیید حذف", "حذف شود؟"): return
        idxs = sorted([int(self.price_tree.item(i, "values")[0]) - 1 for i in sel], reverse=True)
        for idx in idxs:
            if 0 <= idx < len(self.price_data):
                self.price_data = self.price_data.drop(self.price_data.index[idx]).reset_index(drop=True)
        self._refresh_price_table(self.price_search_var.get())


if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()